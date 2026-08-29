"""
ExcelKeywords
=================
Robot Framework library for managing Excel files (.xlsx / .xlsm) using OpenPyXL.
A lightweight alternative to RPA.Excel.Files for basic Excel operations, including:

Open workbook
Create workbook
Read data
Write data
Save workbook
Close workbook

Usage in a .robot file:
    *** Settings ***
    Library    ExcelKeywords.py
"""

import os
import re
from robot.api.deco import keyword, library
from robot.api import logger


@library(scope="TEST")
class ExcelKeywords:
    """A Robot Framework keyword library for handling Excel files using OpenPyXL."""

    def __init__(self):
        self.workbook = None
        self.active_sheet = None
        self.workbook_path = None

    # ------------------------------------------------------------------
    # Open, create, save, and close Excel workbooks (.xlsx / .xlsm)
    # ------------------------------------------------------------------

    @keyword("Open Workbook")
    def open_workbook(self, path, data_only=False):
        """Open an existing Excel file (.xlsx or .xlsm).
        Arguments:

        path : Path to the Excel file.
        data_only : Set to True to read the cached values of formulas instead of the formula expressions.

        Notes:

        .xlsm files are supported automatically.
        Existing macros are preserved when the workbook is saved again.
        """
        import openpyxl

        if not os.path.isfile(path):
            raise FileNotFoundError(f"File Not Found: {path}")

        keep_vba = path.lower().endswith(".xlsm")
        self.workbook = openpyxl.load_workbook(
            path, keep_vba=keep_vba, data_only=data_only
        )
        self.workbook_path = path
        self.active_sheet = self.workbook.active
        logger.info(f"Open file success: {path} (active sheet: {self.active_sheet.title})")
        return self.workbook

    @keyword("Create Workbook")
    def create_workbook(self, path=None, sheet_name="Sheet1"):
        """Create a new workbook in memory.
        The workbook is not saved to a file until Save Workbook is called.

        Arguments:
        path       : File path to save the workbook. You can specify it now or later when calling Save Workbook.
        sheet_name : Name of the first worksheet to create.
        """
        import openpyxl

        self.workbook = openpyxl.Workbook()
        self.workbook.active.title = sheet_name
        self.active_sheet = self.workbook.active
        self.workbook_path = path
        logger.info(f"A new workbook has been created (sheet: {sheet_name})")
        return self.workbook

    @keyword("Close Workbook")
    def close_workbook(self):
        """Close the currently open workbook. Changes are not saved automatically.
        If you want to keep any changes, call Save Workbook before closing.
        """
        if self.workbook:
            self.workbook.close()
            logger.info("Close Workbook")
        self.workbook = None
        self.active_sheet = None
        self.workbook_path = None

    @keyword("Save Workbook")
    def save_workbook(self, path=None):
        """Save the workbook to a file.

        Arguments:
        path : File path to save the workbook.
           If not specified, the workbook will be saved to the same file
           that was opened or created previously.
        """
        self._check_workbook_opened()
        save_path = path or self.workbook_path
        if not save_path:
            raise ValueError("No file path specified for saving the workbook.")

        self.workbook.save(save_path)
        self.workbook_path = save_path
        logger.info(f"Workbook saved successfully: {save_path}")

    # ------------------------------------------------------------------
    # handle worksheet
    # ------------------------------------------------------------------

    @keyword("List Worksheets")
    def list_worksheets(self):
        """Return a list of all worksheet names in the workbook."""
        self._check_workbook_opened()
        return self.workbook.sheetnames

    @keyword("Set Active Worksheet")
    def set_active_worksheet(self, sheet_name):
        """Set the active worksheet.

        Subsequent Read or Write operations will use this worksheet by default
        if sheet_name is not specified.
        """
        self._check_workbook_opened()
        if sheet_name not in self.workbook.sheetnames:
            raise ValueError(f"Worksheet not found '{sheet_name}' (already exists: {self.workbook.sheetnames})")
        self.active_sheet = self.workbook[sheet_name]
        logger.info(f"Set active worksheet to: {sheet_name}")

    @keyword("Create Worksheet")
    def create_worksheet(self, name, content=None, exist_ok=False, header=False):
        """Create a new worksheet in the workbook.

        Arguments:
        name     : Name of the worksheet to create.
        content  : Optional data to populate the worksheet immediately.
               Supported formats:
               - List of lists, for example:
                 [["Name", "Age"], ["Alice", 30]]
               - List of dictionaries, for example:
                 [{"Name": "Alice", "Age": 30}]
        exist_ok : If False (default), an error is raised when the worksheet
               already exists.
               If True, the existing worksheet is used instead of creating
               a new one.
        header   : If True, write a header row when content is a list of
               dictionaries. The keys from the first dictionary are used
               as column headers automatically.
        """
        self._check_workbook_opened()

        if name in self.workbook.sheetnames:
            if not exist_ok:
                raise ValueError(f"Worksheet '{name}' already exists. Set exist_ok=True to use the existing worksheet.")
            sheet = self.workbook[name]
        else:
            sheet = self.workbook.create_sheet(title=name)

        self.active_sheet = sheet

        if content:
            if header and isinstance(content[0], dict):
                keys = list(content[0].keys())
                sheet.append(keys)
                for row in content:
                    sheet.append([row.get(k) for k in keys])
            else:
                for row in content:
                    sheet.append(list(row.values()) if isinstance(row, dict) else list(row))

        logger.info(f"Worksheet '{name}' created successfully.")
        return sheet

    @keyword("Remove Worksheet")
    def remove_worksheet(self, name=None):
        """Delete a worksheet from the workbook.

        Arguments:
        name : Name of the worksheet to delete.
           If not specified, the currently active worksheet is deleted.
        """
        self._check_workbook_opened()
        sheet_name = name or (self.active_sheet.title if self.active_sheet else None)

        if not sheet_name:
            raise ValueError("No worksheet name was provided, and there is no active worksheet.")
        if sheet_name not in self.workbook.sheetnames:
            raise ValueError(f"Worksheet '{sheet_name}' not found. Available worksheets: {self.workbook.sheetnames}")
        if len(self.workbook.sheetnames) == 1:
            raise ValueError("Cannot delete the worksheet because it is the last worksheet in the workbook.")

        self.workbook.remove(self.workbook[sheet_name])
        logger.info(f"ลบ worksheet '{sheet_name}' แล้ว")

        # If the deleted worksheet was the active worksheet,
        # set the first remaining worksheet as the active worksheet.
        if self.active_sheet is None or self.active_sheet.title == sheet_name:
            self.active_sheet = self.workbook.active

    @keyword("Worksheet Exists")
    def worksheet_exists(self, name):
        """
        Return True if a worksheet with the specified name exists in the workbook;
        otherwise, return False.
        """
        self._check_workbook_opened()
        return name in self.workbook.sheetnames

    @keyword("Get Active Worksheet")
    def get_active_worksheet(self):
        """Get the name of the currently active worksheet.

        Returns:
        The name of the active worksheet.
        """

        self._check_workbook_opened()
        return self.active_sheet.title if self.active_sheet else None

    @keyword("Rename Worksheet")
    def rename_worksheet(self, new_name, name=None):
        """Rename a worksheet.

        Arguments:
        new_name : New name for the worksheet.
        name     : Current worksheet name to rename.
                   If not specified, the currently active worksheet is used.
        """
        self._check_workbook_opened()
        sheet = self.workbook[name] if name else self.active_sheet
        if sheet is None:
            raise ValueError("No worksheet available to rename. No worksheet name was provided, and there is no active worksheet.")
        if new_name in self.workbook.sheetnames:
            raise ValueError(f"Worksheet '{new_name}' already exists.")

        old_name = sheet.title
        sheet.title = new_name
        logger.info(f"Worksheet '{old_name}' has been renamed to '{new_name}'.")

    # ------------------------------------------------------------------
    # Read / Write Data
    # ------------------------------------------------------------------

    @keyword("Read Worksheet")
    def read_worksheet(self, sheet_name=None, header=False, start_row=1):
        """Read all data from a worksheet and always return it as a list of dictionaries.

        Arguments:
        sheet_name : Worksheet to read from.
                     If not specified, the currently active worksheet is used.
        header     : If False (default), column letters are used as keys
                     (for example: "A", "B", "C").
                     If True, the first row is used as the dictionary keys
                     (actual column names) and is excluded from the returned data.
        start_row  : Row number to start reading from.
                     Default is row 1.
        """
        from openpyxl.utils import get_column_letter

        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet

        rows = list(sheet.iter_rows(min_row=start_row, values_only=True))
        if not rows:
            return []

        if header:
            keys = rows[0]
            data_rows = rows[1:]
        else:
            keys = [get_column_letter(i + 1) for i in range(len(rows[0]))]
            data_rows = rows

        return [dict(zip(keys, row)) for row in data_rows]

    @keyword("Get Cell Value")
    def get_cell_value(self, row, column, sheet_name=None):
        """Read the value of a cell by specifying its row and column.Rows and columns are 1-based.

        Arguments:
        row        : Row number, for example 1.
        column     : Column identifier. You can use either a column number
                     (1, 2, 3, ...) or a column letter ("A", "B", "C", ...).
        sheet_name : Worksheet to read from.
                     If not specified, the currently active worksheet is used.

        Examples:
            Get Cell Value    row=1    column=1
            Get Cell Value    row=1    column=A
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        col_index = self._resolve_column(column)
        return sheet.cell(row=int(row), column=col_index).value

    @keyword("Set Cell Value")
    def set_cell_value(self, row, column, value, sheet_name=None):
        """Write a value to a cell by specifying its row and column.

        Rows and columns are 1-based.

        Arguments:
        row        : Row number, for example 1.
        column     : Column identifier. You can use either a column number
                 (1, 2, 3, ...) or a column letter ("A", "B", "C", ...).
        value      : Value to write to the cell.
        sheet_name : Worksheet to write to.
                 If not specified, the currently active worksheet is used.

        Examples:
            Set Cell Value    row=1    column=1    value=Hello
            Set Cell Value    row=1    column=A    value=Hello
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        col_index = self._resolve_column(column)
        sheet.cell(row=int(row), column=col_index, value=value)
        logger.info(
            f"Setting cell (row={row}, column={column}) = {value} (sheet: {sheet.title})"
        )

    @keyword("Append Rows To Worksheet")
    def append_rows_to_worksheet(self, rows, sheet_name=None):
        """
        Append multiple rows to the end of a worksheet.

        Arguments:
            rows : List of rows to append, for example:
                [[1, "A"], [2, "B"]]
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        for row in rows:
            sheet.append(row)
        logger.info(f"Appended {len(rows)} rows to worksheet '{sheet.title}'.")

    # ------------------------------------------------------------------
    # Manage worksheet structure (rows, columns, and ranges).
    # ------------------------------------------------------------------

    @keyword("Insert Rows Before")
    def insert_rows_before(self, row, count=1, sheet_name=None):
        """
        Insert blank rows before the specified row.

        Arguments:
            row   : Row number where new rows will be inserted.
                    For example, row=3 inserts new rows before the existing row 3.
            count : Number of rows to insert.
                    Default is 1.
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        sheet.insert_rows(int(row), amount=int(count))
        logger.info(f"Inserted {count} row(s) before row {row} in worksheet '{sheet.title}'.")

    @keyword("Insert Rows After")
    def insert_rows_after(self, row, count=1, sheet_name=None):
        """
        Insert blank rows after the specified row.

        Arguments:
            row   : Row number after which new rows will be inserted.
            count : Number of rows to insert.
                    Default is 1.
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        sheet.insert_rows(int(row) + 1, amount=int(count))
        logger.info(f"Inserted {count} row(s) after row {row} in worksheet '{sheet.title}'.")

    @keyword("Delete Rows")
    def delete_rows(self, row, count=1, sheet_name=None):
        """
        Delete rows from a worksheet.

        Arguments:
            row   : Starting row number to delete.
            count : Number of rows to delete.
                    Default is 1.
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        sheet.delete_rows(int(row), amount=int(count))
        logger.info(f"Deleted {count} row(s) starting at row {row} (worksheet: '{sheet.title}').")

    @keyword("Insert Columns Before")
    def insert_columns_before(self, column, count=1, sheet_name=None):
        """
        Insert blank columns before the specified column.

        Arguments:
            column : Column identifier. You can use either a column number
                    (for example, 3) or a column letter (for example, "C").
            count  : Number of columns to insert.
                    Default is 1.
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        col_index = self._resolve_column(column)
        sheet.insert_cols(col_index, amount=int(count))
        logger.info(f"Inserted {count} column(s) before column {column} in worksheet '{sheet.title}'.")

    @keyword("Insert Columns After")
    def insert_columns_after(self, column, count=1, sheet_name=None):
        """
        Insert blank columns after the specified column.

        Arguments:
            column : Column identifier. You can use either a column number
                    (for example, 3) or a column letter (for example, "C").
            count  : Number of columns to insert.
                    Default is 1.
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        col_index = self._resolve_column(column)
        sheet.insert_cols(col_index + 1, amount=int(count))
        logger.info(f"Inserted {count} column(s) after column {column} in worksheet '{sheet.title}'.")

    @keyword("Delete Columns")
    def delete_columns(self, column, count=1, sheet_name=None):
        """
        Delete columns from a worksheet.

        Arguments:
            column : Starting column to delete. You can use either a column number
                    (for example, 3) or a column letter (for example, "C").
            count  : Number of columns to delete.
                    Default is 1.
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        col_index = self._resolve_column(column)
        sheet.delete_cols(col_index, amount=int(count))
        logger.info(f"Deleted {count} column(s) starting at column {column} (worksheet: '{sheet.title}').")

    @keyword("Move Range")
    def move_range(self, cell_range, rows=0, columns=0, translate=True, sheet_name=None):
        """
        Move a range of cells by the specified number of rows and columns.

        Arguments:
            cell_range : Cell range to move, for example: "E2:E10".
            rows       : Number of rows to move.
                        Positive values move cells down.
                        Negative values move cells up.
            columns    : Number of columns to move.
                        Positive values move cells to the right.
                        Negative values move cells to the left.
            translate  : If True (default), formulas are updated automatically
                        to reflect the new cell positions.

        Note:
            Some versions of openpyxl may have issues when moving cells
            with negative row or column offsets.
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        sheet.move_range(cell_range, rows=int(rows), cols=int(columns), translate=bool(translate))
        logger.info(
            f"Moved range {cell_range} (rows={rows}, columns={columns}) in worksheet '{sheet.title}'."
        )

    @keyword("Clear Cell Range")
    def clear_cell_range(self, cell_range, sheet_name=None):
        """
        Clear the value of a single cell or a range of cells.

        Cell formatting and styles are preserved.

        Arguments:
            cell_range : A single cell (for example, "A1") or a cell range
                        (for example, "B2:B50").
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        for cell in self._iter_cells(sheet, cell_range):
            cell.value = None
        logger.info(f"Cleared values in {cell_range} in worksheet '{sheet.title}'.")

    @keyword("Copy Cell Values")
    def copy_cell_values(self, source_range, target_cell, sheet_name=None, target_sheet_name=None):
        """
        Copy values from a source cell range to a target cell.

        The target cell is used as the top-left corner of the destination range.

        Arguments:
            source_range      : Source cell range, for example: "A1:C3".
            target_cell       : Target cell where the copied data will start,
                                for example: "E1".
            sheet_name        : Source worksheet.
                                If not specified, the currently active worksheet is used.
            target_sheet_name : Target worksheet.
                                If not specified, the source worksheet is used.
        """
        from openpyxl.utils.cell import coordinate_from_string
        from openpyxl.utils import column_index_from_string

        self._check_workbook_opened()
        source_sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        target_sheet = (
            self.workbook[target_sheet_name] if target_sheet_name else source_sheet
        )

        source = source_sheet[source_range]
        rows = source if isinstance(source, tuple) else ((source,),)

        col_letter, start_row = coordinate_from_string(target_cell)
        start_col = column_index_from_string(col_letter)

        for r_offset, row in enumerate(rows):
            for c_offset, cell in enumerate(row):
                target_sheet.cell(
                    row=start_row + r_offset,
                    column=start_col + c_offset,
                    value=cell.value,
                )

        logger.info(
            f"Copied values from {source_range} to {target_cell} "
            f"in worksheet '{source_sheet.title}' -> '{target_sheet.title}'."
        )

    @keyword("Find Empty Row")
    def find_empty_row(self, sheet_name=None):
        """
        Find the first empty row in a worksheet.

        An empty row is a row where all cells are blank.

        This is useful for finding the next available row when appending
        new data to the end of an existing worksheet.
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet

        for row in range(1, sheet.max_row + 2):
            values = [cell.value for cell in sheet[row]]
            if not any(v is not None for v in values):
                return row

        return sheet.max_row + 1

    # ------------------------------------------------------------------
    # (Set styles)
    # ------------------------------------------------------------------

    @keyword("Set Styles")
    def set_styles(
        self,
        cell_range,
        sheet_name=None,
        bold=False,
        italic=False,
        strikethrough=False,
        underline=False,
        font_name=None,
        size=None,
        color=None,
        cell_fill=None,
        align_horizontal=None,
        align_vertical=None,
        number_format=None,
        border=False,
    ):
        """Apply styles to a single cell or a range of cells.
        Parameter names are aligned with RPA.Excel.Files.Set Styles.

        Arguments:
                cell_range        : A single cell (for example, "A1") or a cell range
                                    (for example, "A1:D4").
                sheet_name        : Worksheet to apply the styles to.
                                    If not specified, the currently active worksheet is used.
                bold, italic      : Set to True to enable bold or italic text.
                strikethrough     : Set to True to display text with a strikethrough.
                underline         : Set to True to underline the text.
                font_name, size   : Font name and font size.
                color             : Font color. Accepts either a color name
                                    (for example, "red", "lightblue")
                                    or a hex color code (for example, "FF0000").
                cell_fill         : Cell background color. Accepts either a color name
                                    or a hex color code, same as color.
                align_horizontal  : Horizontal alignment:
                                    "left", "center", or "right".
                align_vertical    : Vertical alignment:
                                    "top", "center", or "bottom".
                number_format     : Number display format, for example:
                                    "#,##0.00", "h:mm AM/PM", or "dd/mm/yyyy".
                border            : Set to True to apply a thin border around the cell(s).
                                    This is an additional feature beyond the original library.

            Examples:
                Set Styles    A1:D4    bold=${TRUE}    cell_fill=lightblue
                ...           align_horizontal=center    number_format=h:mm AM/PM

                Set Styles    E2    strikethrough=${TRUE}    color=FF0000
        """
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet

        font_color = self._resolve_color(color)
        fill_color = self._resolve_color(cell_fill)

        font = Font(
            name=font_name,
            size=size,
            bold=bool(bold),
            italic=bool(italic),
            strike=bool(strikethrough),
            underline="single" if underline else None,
            color=font_color,
        )
        fill = (
            PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
            if fill_color
            else None
        )
        alignment = (
            Alignment(horizontal=align_horizontal, vertical=align_vertical)
            if (align_horizontal or align_vertical)
            else None
        )
        cell_border = None
        if border:
            side = Side(style="thin")
            cell_border = Border(left=side, right=side, top=side, bottom=side)

        for cell in self._iter_cells(sheet, cell_range):
            cell.font = font
            if fill:
                cell.fill = fill
            if alignment:
                cell.alignment = alignment
            if cell_border:
                cell.border = cell_border
            if number_format:
                cell.number_format = number_format

        logger.info(f"Applied styles to {cell_range} in worksheet '{sheet.title}'.")

    # ------------------------------------------------------------------
    # Picture
    # ------------------------------------------------------------------

    @keyword("Insert Image To Worksheet")
    def insert_image_to_worksheet(self, cell, image_path, sheet_name=None, width=None, height=None):
        """Insert an image into a worksheet with its top-left corner anchored
        to the specified cell.

        Arguments:
        cell         : Cell reference where the top-left corner of the image
                   will be placed, for example: "A1".
        image_path   : Path to the image file.
        sheet_name   : Worksheet to insert the image into.
                   If not specified, the currently active worksheet is used.
        width,
        height       : Optional image size in pixels.
                   If not specified, the original image size is used.
        """
        from openpyxl.drawing.image import Image

        self._check_workbook_opened()
        if not os.path.isfile(image_path):
            raise FileNotFoundError(f"Image file not found: {image_path}")

        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        image = Image(image_path)
        if width:
            image.width = int(width)
        if height:
            image.height = int(height)

        sheet.add_image(image, cell)
        logger.info(f"Inserted image '{image_path}' at {cell} in worksheet '{sheet.title}'.")

    # ------------------------------------------------------------------
    # internal helper
    # ------------------------------------------------------------------

    def _check_workbook_opened(self):
        if self.workbook is None:
            raise RuntimeError(
                "Workbook is not open. Please call 'Open Workbook' or 'Create Workbook' first."
            )

    @staticmethod
    def _resolve_column(column):
        """Convert a column value to a 1-based column index.

        Supported formats:
            - Integer values or numeric strings, for example: 1, "1".
            - Excel column letters, for example: "A", "b", "AA".
        """
        column_str = str(column).strip()
        if column_str.isdigit():
            return int(column_str)

        from openpyxl.utils import column_index_from_string

        try:
            return column_index_from_string(column_str.upper())
        except ValueError:
            raise ValueError(
                f"column ไม่ถูกต้อง: '{column}' ต้องเป็นเลข (เช่น 1, 2) หรือตัวอักษรคอลัมน์ (เช่น A, B, AA)"
            )

    @staticmethod
    def _iter_cells(sheet, cell_range):
        """
        Return all cell objects from a single cell or a cell range
        as a flattened list.
        """
        target = sheet[cell_range]
        if isinstance(target, tuple):
            return [cell for row in target for cell in row]
        return [target]

    _COLOR_NAMES = {
        "black": "000000",
        "white": "FFFFFF",
        "red": "FF0000",
        "green": "008000",
        "blue": "0000FF",
        "yellow": "FFFF00",
        "orange": "FFA500",
        "purple": "800080",
        "pink": "FFC0CB",
        "cyan": "00FFFF",
        "magenta": "FF00FF",
        "gray": "808080",
        "grey": "808080",
        "lightblue": "ADD8E6",
        "lightgreen": "90EE90",
        "lightyellow": "FFFFE0",
        "lightgray": "D3D3D3",
        "lightgrey": "D3D3D3",
        "darkred": "8B0000",
        "darkblue": "00008B",
        "darkgreen": "006400",
        "brown": "A52A2A",
    }

    @classmethod
    def _resolve_color(cls, color):
        """
        Convert a color value (color name or hex code) to a hex string
        supported by openpyxl.

        Supported formats:
            - Known color names, such as "red" or "lightblue"
            (case-insensitive).
            - Hex color codes, such as "FF0000" or "#FF0000"
            (6-digit or 8-digit format).
            - None (no color will be applied).
        """
        if not color:
            return None

        color_str = str(color).strip().lstrip("#")
        lowered = color_str.lower()
        if lowered in cls._COLOR_NAMES:
            return cls._COLOR_NAMES[lowered]

        if re.fullmatch(r"[0-9A-Fa-f]{6}", color_str) or re.fullmatch(r"[0-9A-Fa-f]{8}", color_str):
            return color_str.upper()

        raise ValueError(
            f"Invalid color: '{color}'. Expected a recognized color name "
            f"(for example, 'red' or 'lightblue') or a hex color code "
            f"(for example, 'FF0000')."
        )
