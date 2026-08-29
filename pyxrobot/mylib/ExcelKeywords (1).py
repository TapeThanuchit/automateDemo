"""
ExcelKeywords.py
=================
Robot Framework library สำหรับจัดการไฟล์ Excel (.xlsx / .xlsm) โดยใช้ openpyxl
เป็นทางเลือกทดแทน RPA.Excel.Files สำหรับงานพื้นฐาน: Open / Create / Read / Write / Save / Close

การใช้งานใน .robot:
    *** Settings ***
    Library    ExcelKeywords.py
"""

import os
from robot.api.deco import keyword, library
from robot.api import logger


@library(scope="TEST")
class ExcelKeywords:
    """Keyword library สำหรับจัดการไฟล์ Excel ด้วย openpyxl"""

    def __init__(self):
        self.workbook = None
        self.active_sheet = None
        self.workbook_path = None

    # ------------------------------------------------------------------
    # เปิด / สร้าง / ปิด / บันทึก workbook
    # ------------------------------------------------------------------

    @keyword("Open Workbook")
    def open_workbook(self, path, data_only=False):
        """เปิดไฟล์ Excel ที่มีอยู่แล้ว (.xlsx หรือ .xlsm)

        - path        : path ของไฟล์ Excel
        - data_only   : True เพื่ออ่านค่าที่ cache ไว้ของสูตร (ไม่ใช่ตัวสูตรเอง)

        รองรับ .xlsm อัตโนมัติ (จะเก็บ macro ไว้เมื่อ save ต่อ)
        """
        import openpyxl

        if not os.path.isfile(path):
            raise FileNotFoundError(f"ไม่พบไฟล์: {path}")

        keep_vba = path.lower().endswith(".xlsm")
        self.workbook = openpyxl.load_workbook(
            path, keep_vba=keep_vba, data_only=data_only
        )
        self.workbook_path = path
        self.active_sheet = self.workbook.active
        logger.info(f"เปิดไฟล์สำเร็จ: {path} (active sheet: {self.active_sheet.title})")
        return self.workbook

    @keyword("Create Workbook")
    def create_workbook(self, path=None, sheet_name="Sheet1"):
        """สร้าง workbook ใหม่ในหน่วยความจำ (ยังไม่บันทึกลงไฟล์จนกว่าจะเรียก Save Workbook)

        - path       : path ที่ต้องการบันทึก (ระบุตอนนี้หรือตอนเรียก Save Workbook ก็ได้)
        - sheet_name : ชื่อ sheet แรกที่จะสร้างให้
        """
        import openpyxl

        self.workbook = openpyxl.Workbook()
        self.workbook.active.title = sheet_name
        self.active_sheet = self.workbook.active
        self.workbook_path = path
        logger.info(f"สร้าง workbook ใหม่แล้ว (sheet: {sheet_name})")
        return self.workbook

    @keyword("Close Workbook")
    def close_workbook(self):
        """ปิด workbook ที่เปิดอยู่ (ไม่ได้ save ให้อัตโนมัติ ต้องเรียก Save Workbook ก่อนถ้าต้องการเก็บการเปลี่ยนแปลง)"""
        if self.workbook:
            self.workbook.close()
            logger.info("ปิด workbook แล้ว")
        self.workbook = None
        self.active_sheet = None
        self.workbook_path = None

    @keyword("Save Workbook")
    def save_workbook(self, path=None):
        """บันทึก workbook ลงไฟล์

        - path : ถ้าไม่ระบุ จะบันทึกทับไฟล์เดิมที่เปิดไว้ (path ตอน Open/Create Workbook)
        """
        self._check_workbook_opened()
        save_path = path or self.workbook_path
        if not save_path:
            raise ValueError("ไม่ได้ระบุ path สำหรับบันทึกไฟล์ (ไม่มี path เดิมและไม่ได้ระบุ path ใหม่)")

        self.workbook.save(save_path)
        self.workbook_path = save_path
        logger.info(f"บันทึกไฟล์แล้ว: {save_path}")

    # ------------------------------------------------------------------
    # จัดการ worksheet
    # ------------------------------------------------------------------

    @keyword("List Worksheets")
    def list_worksheets(self):
        """คืนค่ารายชื่อ worksheet ทั้งหมดใน workbook เป็น list"""
        self._check_workbook_opened()
        return self.workbook.sheetnames

    @keyword("Set Active Worksheet")
    def set_active_worksheet(self, sheet_name):
        """เปลี่ยน worksheet ที่ใช้งานอยู่ (คำสั่ง Read/Write ถัดไปจะทำงานกับ sheet นี้ถ้าไม่ระบุ sheet_name)"""
        self._check_workbook_opened()
        if sheet_name not in self.workbook.sheetnames:
            raise ValueError(f"ไม่พบ worksheet ชื่อ '{sheet_name}' (มีอยู่: {self.workbook.sheetnames})")
        self.active_sheet = self.workbook[sheet_name]
        logger.info(f"เปลี่ยน active worksheet เป็น: {sheet_name}")

    @keyword("Create Worksheet")
    def create_worksheet(self, name, content=None, exist_ok=False, header=False):
        """สร้าง worksheet ใหม่ใน workbook

        - name     : ชื่อ worksheet ที่จะสร้าง
        - content  : (ไม่บังคับ) ข้อมูลที่จะใส่ให้ทันที เป็น list ของ list เช่น
                     [["Name", "Age"], ["Alice", 30]]
                     หรือ list ของ dict เช่น [{"Name": "Alice", "Age": 30}]
        - exist_ok : False (ค่าเริ่มต้น) จะ error ถ้ามีชื่อ sheet นี้อยู่แล้ว
                     True จะใช้ sheet เดิมที่มีอยู่แทนการสร้างใหม่
        - header   : True เพื่อเขียนแถวหัวตาราง (ใช้เมื่อ content เป็น list ของ dict
                     จะเอา key ของ dict แรกมาเป็นหัวตารางให้อัตโนมัติ)
        """
        self._check_workbook_opened()

        if name in self.workbook.sheetnames:
            if not exist_ok:
                raise ValueError(f"มี worksheet ชื่อ '{name}' อยู่แล้ว (ตั้ง exist_ok=True ถ้าต้องการใช้ sheet เดิม)")
            sheet = self.workbook[name]
        else:
            sheet = self.workbook.create_sheet(title=name)

        self.active_sheet = sheet

        if content:
            if header and isinstance(content[0], dict):
                keys = list(content[0].keys())
                sheet.append(keys)
                for row in content:
                    sheet.append([row.get(k) for k in keys])
            else:
                for row in content:
                    sheet.append(list(row.values()) if isinstance(row, dict) else list(row))

        logger.info(f"สร้าง worksheet '{name}' แล้ว")
        return sheet

    @keyword("Remove Worksheet")
    def remove_worksheet(self, name=None):
        """ลบ worksheet ออกจาก workbook

        - name : ชื่อ worksheet ที่จะลบ (ไม่ระบุ = ลบ active worksheet ปัจจุบัน)
        """
        self._check_workbook_opened()
        sheet_name = name or (self.active_sheet.title if self.active_sheet else None)

        if not sheet_name:
            raise ValueError("ไม่ได้ระบุชื่อ worksheet ที่จะลบ และไม่มี active worksheet อยู่")
        if sheet_name not in self.workbook.sheetnames:
            raise ValueError(f"ไม่พบ worksheet ชื่อ '{sheet_name}' (มีอยู่: {self.workbook.sheetnames})")
        if len(self.workbook.sheetnames) == 1:
            raise ValueError("ไม่สามารถลบได้ เพราะเป็น worksheet สุดท้ายของ workbook")

        self.workbook.remove(self.workbook[sheet_name])
        logger.info(f"ลบ worksheet '{sheet_name}' แล้ว")

        # ถ้า sheet ที่ลบคือ active sheet เดิม ให้ย้าย active ไปที่ sheet แรกที่เหลือ
        if self.active_sheet is None or self.active_sheet.title == sheet_name:
            self.active_sheet = self.workbook.active

    # ------------------------------------------------------------------
    # อ่าน / เขียนข้อมูล
    # ------------------------------------------------------------------

    @keyword("Read Worksheet")
    def read_worksheet(self, sheet_name=None, header=False, start_row=1):
        """อ่านข้อมูลทั้ง worksheet

        - sheet_name : ระบุ sheet ที่จะอ่าน (ไม่ระบุ = ใช้ active worksheet)
        - header     : True เพื่อใช้แถวแรกเป็นชื่อคอลัมน์ และคืนค่าเป็น list ของ dict
                       False เพื่อคืนค่าเป็น list ของ list ตรงๆ (ค่าเริ่มต้น)
        - start_row  : แถวที่เริ่มอ่าน (ค่าเริ่มต้น = แถวที่ 1)
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet

        rows = list(sheet.iter_rows(min_row=start_row, values_only=True))
        if not rows:
            return []

        if header:
            headers = rows[0]
            return [dict(zip(headers, row)) for row in rows[1:]]

        return [list(row) for row in rows]

    @keyword("Get Cell Value")
    def get_cell_value(self, row, column, sheet_name=None):
        """อ่านค่าจาก cell โดยระบุตำแหน่งเป็น row/column (เริ่มนับที่ 1)

        - row        : เลขแถว เช่น 1
        - column     : คอลัมน์ ระบุเป็นเลข เช่น 1, 2, 3 หรือตัวอักษร เช่น "A", "B", "C" ก็ได้
        - sheet_name : sheet ที่จะอ่าน (ไม่ระบุ = ใช้ active worksheet)

        ตัวอย่าง:
            Get Cell Value    row=1    column=1
            Get Cell Value    row=1    column=A
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        col_index = self._resolve_column(column)
        return sheet.cell(row=int(row), column=col_index).value

    @keyword("Set Cell Value")
    def set_cell_value(self, row, column, value, sheet_name=None):
        """เขียนค่าลง cell โดยระบุตำแหน่งเป็น row/column (เริ่มนับที่ 1)

        - row        : เลขแถว เช่น 1
        - column     : คอลัมน์ ระบุเป็นเลข เช่น 1, 2, 3 หรือตัวอักษร เช่น "A", "B", "C" ก็ได้
        - value      : ค่าที่จะเขียน
        - sheet_name : sheet ที่จะเขียน (ไม่ระบุ = ใช้ active worksheet)

        ตัวอย่าง:
            Set Cell Value    row=1    column=1    value=Hello
            Set Cell Value    row=1    column=A    value=Hello
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        col_index = self._resolve_column(column)
        sheet.cell(row=int(row), column=col_index, value=value)
        logger.info(
            f"ตั้งค่า cell (row={row}, column={column}) = {value} (sheet: {sheet.title})"
        )

    @keyword("Append Rows To Worksheet")
    def append_rows_to_worksheet(self, rows, sheet_name=None):
        """เพิ่มข้อมูลหลายแถวต่อท้าย worksheet

        - rows : list ของ list เช่น [[1, "A"], [2, "B"]]
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        for row in rows:
            sheet.append(row)
        logger.info(f"เพิ่ม {len(rows)} แถวใน sheet: {sheet.title}")

    # ------------------------------------------------------------------
    # จัดรูปแบบ (styles)
    # ------------------------------------------------------------------

    @keyword("Set Styles")
    def set_styles(
        self,
        cell_range,
        sheet_name=None,
        font_name=None,
        font_size=None,
        bold=False,
        italic=False,
        font_color=None,
        bg_color=None,
        horizontal_alignment=None,
        vertical_alignment=None,
        border=False,
    ):
        """ตั้งค่ารูปแบบ (style) ให้กับ cell เดียวหรือช่วง cell

        - cell_range            : cell เดียว เช่น "A1" หรือช่วง เช่น "A1:C5"
        - sheet_name            : sheet ที่จะตั้งค่า (ไม่ระบุ = active worksheet)
        - font_name, font_size  : ชื่อและขนาดตัวอักษร
        - bold, italic          : True/False
        - font_color            : สีตัวอักษร เป็น hex ไม่ต้องมี # เช่น "FF0000"
        - bg_color               : สีพื้นหลัง cell เป็น hex เช่น "FFFF00"
        - horizontal_alignment  : "left" / "center" / "right"
        - vertical_alignment    : "top" / "center" / "bottom"
        - border                : True เพื่อใส่เส้นขอบบางรอบ cell

        ตัวอย่าง:
            Set Styles    A1:C1    bold=${TRUE}    bg_color=FFFF00    horizontal_alignment=center
        """
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet

        font = Font(
            name=font_name,
            size=font_size,
            bold=bool(bold),
            italic=bool(italic),
            color=font_color,
        )
        fill = (
            PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
            if bg_color
            else None
        )
        alignment = (
            Alignment(horizontal=horizontal_alignment, vertical=vertical_alignment)
            if (horizontal_alignment or vertical_alignment)
            else None
        )
        cell_border = None
        if border:
            side = Side(style="thin")
            cell_border = Border(left=side, right=side, top=side, bottom=side)

        target = sheet[cell_range]
        # cell เดียวจะได้ Cell object ตรงๆ / ช่วง cell จะได้ tuple ของ tuple ของ Cell
        cells = [target] if not isinstance(target, tuple) else [c for row in target for c in row]

        for cell in cells:
            cell.font = font
            if fill:
                cell.fill = fill
            if alignment:
                cell.alignment = alignment
            if cell_border:
                cell.border = cell_border

        logger.info(f"ตั้งค่า style ให้ {cell_range} แล้ว (sheet: {sheet.title})")

    # ------------------------------------------------------------------
    # internal helper
    # ------------------------------------------------------------------

    def _check_workbook_opened(self):
        if self.workbook is None:
            raise RuntimeError(
                "ยังไม่ได้เปิดหรือสร้าง workbook กรุณาเรียก 'Open Workbook' หรือ 'Create Workbook' ก่อน"
            )

    @staticmethod
    def _resolve_column(column):
        """แปลงค่า column ที่รับมาให้เป็นเลข index (1-based)

        รองรับทั้ง:
        - เลข หรือ string ตัวเลข เช่น 1, "1"
        - ตัวอักษรคอลัมน์ เช่น "A", "b", "AA"
        """
        column_str = str(column).strip()
        if column_str.isdigit():
            return int(column_str)

        from openpyxl.utils import column_index_from_string

        try:
            return column_index_from_string(column_str.upper())
        except ValueError:
            raise ValueError(
                f"column ไม่ถูกต้อง: '{column}' ต้องเป็นเลข (เช่น 1, 2) หรือตัวอักษรคอลัมน์ (เช่น A, B, AA)"
            )
