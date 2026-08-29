"""
ExcelKeywords.py
=================
Robot Framework library สำหรับจัดการไฟล์ Excel (.xlsx / .xlsm) โดยใช้ openpyxl
เป็นทางเลือกทดแทน RPA.Excel.Files สำหรับงานพื้นฐาน: Open / Create / Read / Write / Save / Close

การใช้งานใน .robot:
    *** Settings ***
    Library    ExcelKeywords.py
"""

import os
import re
from robot.api.deco import keyword, library
from robot.api import logger


@library(scope="TEST")
class ExcelKeywords:
    """Keyword library สำหรับจัดการไฟล์ Excel ด้วย openpyxl"""

    def __init__(self):
        self.workbook = None
        self.active_sheet = None
        self.workbook_path = None

    # ------------------------------------------------------------------
    # เปิด / สร้าง / ปิด / บันทึก workbook
    # ------------------------------------------------------------------

    @keyword("Open Workbook")
    def open_workbook(self, path, data_only=False):
        """เปิดไฟล์ Excel ที่มีอยู่แล้ว (.xlsx หรือ .xlsm)

        - path        : path ของไฟล์ Excel
        - data_only   : True เพื่ออ่านค่าที่ cache ไว้ของสูตร (ไม่ใช่ตัวสูตรเอง)

        รองรับ .xlsm อัตโนมัติ (จะเก็บ macro ไว้เมื่อ save ต่อ)
        """
        import openpyxl

        if not os.path.isfile(path):
            raise FileNotFoundError(f"ไม่พบไฟล์: {path}")

        keep_vba = path.lower().endswith(".xlsm")
        self.workbook = openpyxl.load_workbook(
            path, keep_vba=keep_vba, data_only=data_only
        )
        self.workbook_path = path
        self.active_sheet = self.workbook.active
        logger.info(f"เปิดไฟล์สำเร็จ: {path} (active sheet: {self.active_sheet.title})")
        return self.workbook

    @keyword("Create Workbook")
    def create_workbook(self, path=None, sheet_name="Sheet1"):
        """สร้าง workbook ใหม่ในหน่วยความจำ (ยังไม่บันทึกลงไฟล์จนกว่าจะเรียก Save Workbook)

        - path       : path ที่ต้องการบันทึก (ระบุตอนนี้หรือตอนเรียก Save Workbook ก็ได้)
        - sheet_name : ชื่อ sheet แรกที่จะสร้างให้
        """
        import openpyxl

        self.workbook = openpyxl.Workbook()
        self.workbook.active.title = sheet_name
        self.active_sheet = self.workbook.active
        self.workbook_path = path
        logger.info(f"สร้าง workbook ใหม่แล้ว (sheet: {sheet_name})")
        return self.workbook

    @keyword("Close Workbook")
    def close_workbook(self):
        """ปิด workbook ที่เปิดอยู่ (ไม่ได้ save ให้อัตโนมัติ ต้องเรียก Save Workbook ก่อนถ้าต้องการเก็บการเปลี่ยนแปลง)"""
        if self.workbook:
            self.workbook.close()
            logger.info("ปิด workbook แล้ว")
        self.workbook = None
        self.active_sheet = None
        self.workbook_path = None

    @keyword("Save Workbook")
    def save_workbook(self, path=None):
        """บันทึก workbook ลงไฟล์

        - path : ถ้าไม่ระบุ จะบันทึกทับไฟล์เดิมที่เปิดไว้ (path ตอน Open/Create Workbook)
        """
        self._check_workbook_opened()
        save_path = path or self.workbook_path
        if not save_path:
            raise ValueError("ไม่ได้ระบุ path สำหรับบันทึกไฟล์ (ไม่มี path เดิมและไม่ได้ระบุ path ใหม่)")

        self.workbook.save(save_path)
        self.workbook_path = save_path
        logger.info(f"บันทึกไฟล์แล้ว: {save_path}")

    # ------------------------------------------------------------------
    # จัดการ worksheet
    # ------------------------------------------------------------------

    @keyword("List Worksheets")
    def list_worksheets(self):
        """คืนค่ารายชื่อ worksheet ทั้งหมดใน workbook เป็น list"""
        self._check_workbook_opened()
        return self.workbook.sheetnames

    @keyword("Set Active Worksheet")
    def set_active_worksheet(self, sheet_name):
        """เปลี่ยน worksheet ที่ใช้งานอยู่ (คำสั่ง Read/Write ถัดไปจะทำงานกับ sheet นี้ถ้าไม่ระบุ sheet_name)"""
        self._check_workbook_opened()
        if sheet_name not in self.workbook.sheetnames:
            raise ValueError(f"ไม่พบ worksheet ชื่อ '{sheet_name}' (มีอยู่: {self.workbook.sheetnames})")
        self.active_sheet = self.workbook[sheet_name]
        logger.info(f"เปลี่ยน active worksheet เป็น: {sheet_name}")

    @keyword("Create Worksheet")
    def create_worksheet(self, name, content=None, exist_ok=False, header=False):
        """สร้าง worksheet ใหม่ใน workbook

        - name     : ชื่อ worksheet ที่จะสร้าง
        - content  : (ไม่บังคับ) ข้อมูลที่จะใส่ให้ทันที เป็น list ของ list เช่น
                     [["Name", "Age"], ["Alice", 30]]
                     หรือ list ของ dict เช่น [{"Name": "Alice", "Age": 30}]
        - exist_ok : False (ค่าเริ่มต้น) จะ error ถ้ามีชื่อ sheet นี้อยู่แล้ว
                     True จะใช้ sheet เดิมที่มีอยู่แทนการสร้างใหม่
        - header   : True เพื่อเขียนแถวหัวตาราง (ใช้เมื่อ content เป็น list ของ dict
                     จะเอา key ของ dict แรกมาเป็นหัวตารางให้อัตโนมัติ)
        """
        self._check_workbook_opened()

        if name in self.workbook.sheetnames:
            if not exist_ok:
                raise ValueError(f"มี worksheet ชื่อ '{name}' อยู่แล้ว (ตั้ง exist_ok=True ถ้าต้องการใช้ sheet เดิม)")
            sheet = self.workbook[name]
        else:
            sheet = self.workbook.create_sheet(title=name)

        self.active_sheet = sheet

        if content:
            if header and isinstance(content[0], dict):
                keys = list(content[0].keys())
                sheet.append(keys)
                for row in content:
                    sheet.append([row.get(k) for k in keys])
            else:
                for row in content:
                    sheet.append(list(row.values()) if isinstance(row, dict) else list(row))

        logger.info(f"สร้าง worksheet '{name}' แล้ว")
        return sheet

    @keyword("Remove Worksheet")
    def remove_worksheet(self, name=None):
        """ลบ worksheet ออกจาก workbook

        - name : ชื่อ worksheet ที่จะลบ (ไม่ระบุ = ลบ active worksheet ปัจจุบัน)
        """
        self._check_workbook_opened()
        sheet_name = name or (self.active_sheet.title if self.active_sheet else None)

        if not sheet_name:
            raise ValueError("ไม่ได้ระบุชื่อ worksheet ที่จะลบ และไม่มี active worksheet อยู่")
        if sheet_name not in self.workbook.sheetnames:
            raise ValueError(f"ไม่พบ worksheet ชื่อ '{sheet_name}' (มีอยู่: {self.workbook.sheetnames})")
        if len(self.workbook.sheetnames) == 1:
            raise ValueError("ไม่สามารถลบได้ เพราะเป็น worksheet สุดท้ายของ workbook")

        self.workbook.remove(self.workbook[sheet_name])
        logger.info(f"ลบ worksheet '{sheet_name}' แล้ว")

        # ถ้า sheet ที่ลบคือ active sheet เดิม ให้ย้าย active ไปที่ sheet แรกที่เหลือ
        if self.active_sheet is None or self.active_sheet.title == sheet_name:
            self.active_sheet = self.workbook.active

    @keyword("Worksheet Exists")
    def worksheet_exists(self, name):
        """คืนค่า True/False ว่ามี worksheet ชื่อนี้อยู่ใน workbook หรือไม่"""
        self._check_workbook_opened()
        return name in self.workbook.sheetnames

    @keyword("Get Active Worksheet")
    def get_active_worksheet(self):
        """คืนค่าชื่อ worksheet ที่ active อยู่ในปัจจุบัน"""
        self._check_workbook_opened()
        return self.active_sheet.title if self.active_sheet else None

    @keyword("Rename Worksheet")
    def rename_worksheet(self, new_name, name=None):
        """เปลี่ยนชื่อ worksheet

        - new_name : ชื่อใหม่ที่จะตั้ง
        - name     : ชื่อ worksheet เดิมที่จะเปลี่ยน (ไม่ระบุ = ใช้ active worksheet)
        """
        self._check_workbook_opened()
        sheet = self.workbook[name] if name else self.active_sheet
        if sheet is None:
            raise ValueError("ไม่มี worksheet ให้เปลี่ยนชื่อ (ไม่ได้ระบุ name และไม่มี active worksheet)")
        if new_name in self.workbook.sheetnames:
            raise ValueError(f"มี worksheet ชื่อ '{new_name}' อยู่แล้ว")

        old_name = sheet.title
        sheet.title = new_name
        logger.info(f"เปลี่ยนชื่อ worksheet '{old_name}' เป็น '{new_name}'")

    # ------------------------------------------------------------------
    # อ่าน / เขียนข้อมูล
    # ------------------------------------------------------------------

    @keyword("Read Worksheet")
    def read_worksheet(self, sheet_name=None, header=False, start_row=1):
        """อ่านข้อมูลทั้ง worksheet คืนค่าเป็น list ของ dict เสมอ

        - sheet_name : ระบุ sheet ที่จะอ่าน (ไม่ระบุ = ใช้ active worksheet)
        - header     : False (ค่าเริ่มต้น) จะใช้ตัวอักษรคอลัมน์เป็น key เช่น "A", "B", "C"
                       True จะใช้แถวแรกเป็น key (ชื่อคอลัมน์จริง) และไม่รวมแถวแรกในผลลัพธ์
        - start_row  : แถวที่เริ่มอ่าน (ค่าเริ่มต้น = แถวที่ 1)
        """
        from openpyxl.utils import get_column_letter

        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet

        rows = list(sheet.iter_rows(min_row=start_row, values_only=True))
        if not rows:
            return []

        if header:
            keys = rows[0]
            data_rows = rows[1:]
        else:
            keys = [get_column_letter(i + 1) for i in range(len(rows[0]))]
            data_rows = rows

        return [dict(zip(keys, row)) for row in data_rows]

    # @keyword("Read Worksheet")
    # def read_worksheet(self, sheet_name=None, header=False, start_row=1):
    #     """อ่านข้อมูลทั้ง worksheet

    #     - sheet_name : ระบุ sheet ที่จะอ่าน (ไม่ระบุ = ใช้ active worksheet)
    #     - header     : True เพื่อใช้แถวแรกเป็นชื่อคอลัมน์ และคืนค่าเป็น list ของ dict
    #                    False เพื่อคืนค่าเป็น list ของ list ตรงๆ (ค่าเริ่มต้น)
    #     - start_row  : แถวที่เริ่มอ่าน (ค่าเริ่มต้น = แถวที่ 1)
    #     """
    #     self._check_workbook_opened()
    #     sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet

    #     rows = list(sheet.iter_rows(min_row=start_row, values_only=True))
    #     if not rows:
    #         return []

    #     if header:
    #         headers = rows[0]
    #         return [dict(zip(headers, row)) for row in rows[1:]]

    #     return [list(row) for row in rows]

    @keyword("Get Cell Value")
    def get_cell_value(self, row, column, sheet_name=None):
        """อ่านค่าจาก cell โดยระบุตำแหน่งเป็น row/column (เริ่มนับที่ 1)

        - row        : เลขแถว เช่น 1
        - column     : คอลัมน์ ระบุเป็นเลข เช่น 1, 2, 3 หรือตัวอักษร เช่น "A", "B", "C" ก็ได้
        - sheet_name : sheet ที่จะอ่าน (ไม่ระบุ = ใช้ active worksheet)

        ตัวอย่าง:
            Get Cell Value    row=1    column=1
            Get Cell Value    row=1    column=A
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        col_index = self._resolve_column(column)
        return sheet.cell(row=int(row), column=col_index).value

    @keyword("Set Cell Value")
    def set_cell_value(self, row, column, value, sheet_name=None):
        """เขียนค่าลง cell โดยระบุตำแหน่งเป็น row/column (เริ่มนับที่ 1)

        - row        : เลขแถว เช่น 1
        - column     : คอลัมน์ ระบุเป็นเลข เช่น 1, 2, 3 หรือตัวอักษร เช่น "A", "B", "C" ก็ได้
        - value      : ค่าที่จะเขียน
        - sheet_name : sheet ที่จะเขียน (ไม่ระบุ = ใช้ active worksheet)

        ตัวอย่าง:
            Set Cell Value    row=1    column=1    value=Hello
            Set Cell Value    row=1    column=A    value=Hello
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        col_index = self._resolve_column(column)
        sheet.cell(row=int(row), column=col_index, value=value)
        logger.info(
            f"ตั้งค่า cell (row={row}, column={column}) = {value} (sheet: {sheet.title})"
        )

    @keyword("Append Rows To Worksheet")
    def append_rows_to_worksheet(self, rows, sheet_name=None):
        """เพิ่มข้อมูลหลายแถวต่อท้าย worksheet

        - rows : list ของ list เช่น [[1, "A"], [2, "B"]]
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        for row in rows:
            sheet.append(row)
        logger.info(f"เพิ่ม {len(rows)} แถวใน sheet: {sheet.title}")

    # ------------------------------------------------------------------
    # จัดการโครงสร้าง (rows / columns / range)
    # ------------------------------------------------------------------

    @keyword("Insert Rows Before")
    def insert_rows_before(self, row, count=1, sheet_name=None):
        """แทรกแถวว่างก่อนแถวที่ระบุ

        - row   : เลขแถวที่จะแทรกไว้ข้างหน้า (เช่น 3 = แทรกแถวใหม่ก่อนแถวที่ 3 เดิม)
        - count : จำนวนแถวที่จะแทรก (ค่าเริ่มต้น 1)
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        sheet.insert_rows(int(row), amount=int(count))
        logger.info(f"แทรก {count} แถวก่อนแถวที่ {row} (sheet: {sheet.title})")

    @keyword("Insert Rows After")
    def insert_rows_after(self, row, count=1, sheet_name=None):
        """แทรกแถวว่างหลังแถวที่ระบุ

        - row   : เลขแถวที่จะแทรกไว้ข้างหลัง
        - count : จำนวนแถวที่จะแทรก (ค่าเริ่มต้น 1)
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        sheet.insert_rows(int(row) + 1, amount=int(count))
        logger.info(f"แทรก {count} แถวหลังแถวที่ {row} (sheet: {sheet.title})")

    @keyword("Delete Rows")
    def delete_rows(self, row, count=1, sheet_name=None):
        """ลบแถวออกจาก worksheet

        - row   : เลขแถวที่จะเริ่มลบ
        - count : จำนวนแถวที่จะลบ (ค่าเริ่มต้น 1)
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        sheet.delete_rows(int(row), amount=int(count))
        logger.info(f"ลบ {count} แถว เริ่มจากแถวที่ {row} (sheet: {sheet.title})")

    @keyword("Insert Columns Before")
    def insert_columns_before(self, column, count=1, sheet_name=None):
        """แทรกคอลัมน์ว่างก่อนคอลัมน์ที่ระบุ

        - column : คอลัมน์ (เลขหรือตัวอักษร เช่น 3 หรือ "C")
        - count  : จำนวนคอลัมน์ที่จะแทรก (ค่าเริ่มต้น 1)
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        col_index = self._resolve_column(column)
        sheet.insert_cols(col_index, amount=int(count))
        logger.info(f"แทรก {count} คอลัมน์ก่อนคอลัมน์ {column} (sheet: {sheet.title})")

    @keyword("Insert Columns After")
    def insert_columns_after(self, column, count=1, sheet_name=None):
        """แทรกคอลัมน์ว่างหลังคอลัมน์ที่ระบุ

        - column : คอลัมน์ (เลขหรือตัวอักษร เช่น 3 หรือ "C")
        - count  : จำนวนคอลัมน์ที่จะแทรก (ค่าเริ่มต้น 1)
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        col_index = self._resolve_column(column)
        sheet.insert_cols(col_index + 1, amount=int(count))
        logger.info(f"แทรก {count} คอลัมน์หลังคอลัมน์ {column} (sheet: {sheet.title})")

    @keyword("Delete Columns")
    def delete_columns(self, column, count=1, sheet_name=None):
        """ลบคอลัมน์ออกจาก worksheet

        - column : คอลัมน์ (เลขหรือตัวอักษร เช่น 3 หรือ "C") ที่จะเริ่มลบ
        - count  : จำนวนคอลัมน์ที่จะลบ (ค่าเริ่มต้น 1)
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        col_index = self._resolve_column(column)
        sheet.delete_cols(col_index, amount=int(count))
        logger.info(f"ลบ {count} คอลัมน์ เริ่มจากคอลัมน์ {column} (sheet: {sheet.title})")

    @keyword("Move Range")
    def move_range(self, cell_range, rows=0, columns=0, translate=True, sheet_name=None):
        """ย้ายช่วง cell ไปตามจำนวนแถว/คอลัมน์ที่กำหนด

        - cell_range : ช่วง cell เช่น "E2:E10"
        - rows       : จำนวนแถวที่จะย้าย (ค่าบวก = ย้ายลง, ค่าลบ = ย้ายขึ้น)
        - columns    : จำนวนคอลัมน์ที่จะย้าย (ค่าบวก = ย้ายขวา, ค่าลบ = ย้ายซ้าย)
        - translate  : True (ค่าเริ่มต้น) เพื่อแปลสูตรให้อ้างอิงตำแหน่งใหม่ตามไปด้วย

        หมายเหตุ: openpyxl มีบั๊กเรื่องการย้ายด้วยค่า rows/columns ติดลบในบางเวอร์ชัน
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        sheet.move_range(cell_range, rows=int(rows), cols=int(columns), translate=bool(translate))
        logger.info(
            f"ย้ายช่วง {cell_range} (rows={rows}, columns={columns}) แล้ว (sheet: {sheet.title})"
        )

    @keyword("Clear Cell Range")
    def clear_cell_range(self, cell_range, sheet_name=None):
        """ล้างค่าใน cell เดียวหรือช่วง cell (ไม่ลบ format)

        - cell_range : cell เดียว เช่น "A1" หรือช่วง เช่น "B2:B50"
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        for cell in self._iter_cells(sheet, cell_range):
            cell.value = None
        logger.info(f"ล้างค่าในช่วง {cell_range} แล้ว (sheet: {sheet.title})")

    @keyword("Copy Cell Values")
    def copy_cell_values(self, source_range, target_cell, sheet_name=None, target_sheet_name=None):
        """คัดลอกค่าจากช่วง cell ต้นทางไปวางที่ cell ปลายทาง (มุมบนซ้าย)

        - source_range      : ช่วง cell ต้นทาง เช่น "A1:C3"
        - target_cell        : cell ปลายทางที่จะเริ่มวาง เช่น "E1"
        - sheet_name         : sheet ต้นทาง (ไม่ระบุ = active worksheet)
        - target_sheet_name  : sheet ปลายทาง (ไม่ระบุ = sheet เดียวกับต้นทาง)
        """
        from openpyxl.utils.cell import coordinate_from_string
        from openpyxl.utils import column_index_from_string

        self._check_workbook_opened()
        source_sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        target_sheet = (
            self.workbook[target_sheet_name] if target_sheet_name else source_sheet
        )

        source = source_sheet[source_range]
        rows = source if isinstance(source, tuple) else ((source,),)

        col_letter, start_row = coordinate_from_string(target_cell)
        start_col = column_index_from_string(col_letter)

        for r_offset, row in enumerate(rows):
            for c_offset, cell in enumerate(row):
                target_sheet.cell(
                    row=start_row + r_offset,
                    column=start_col + c_offset,
                    value=cell.value,
                )

        logger.info(
            f"คัดลอกค่าจาก {source_range} ไปที่ {target_cell} แล้ว "
            f"(sheet: {source_sheet.title} -> {target_sheet.title})"
        )

    @keyword("Find Empty Row")
    def find_empty_row(self, sheet_name=None):
        """หาเลขแถวว่างแถวแรก (แถวที่ทุก cell ไม่มีค่าเลย) ใน worksheet

        เหมาะสำหรับหาตำแหน่งที่จะเขียนข้อมูลแถวใหม่ต่อท้ายข้อมูลเดิม
        """
        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet

        for row in range(1, sheet.max_row + 2):
            values = [cell.value for cell in sheet[row]]
            if not any(v is not None for v in values):
                return row

        return sheet.max_row + 1

    # ------------------------------------------------------------------
    # จัดรูปแบบ (styles)
    # ------------------------------------------------------------------

    @keyword("Set Styles")
    def set_styles(
        self,
        cell_range,
        sheet_name=None,
        bold=False,
        italic=False,
        strikethrough=False,
        underline=False,
        font_name=None,
        size=None,
        color=None,
        cell_fill=None,
        align_horizontal=None,
        align_vertical=None,
        number_format=None,
        border=False,
    ):
        """ตั้งค่ารูปแบบ (style) ให้กับ cell เดียวหรือช่วง cell
        (ชื่อ parameter ตรงกับ RPA.Excel.Files.Set Styles)

        - cell_range           : cell เดียว เช่น "A1" หรือช่วง เช่น "A1:D4"
        - sheet_name            : sheet ที่จะตั้งค่า (ไม่ระบุ = active worksheet)
        - bold, italic          : True/False
        - strikethrough         : True เพื่อขีดฆ่าตัวอักษร
        - underline             : True เพื่อขีดเส้นใต้
        - font_name, size       : ชื่อและขนาดตัวอักษร
        - color                 : สีตัวอักษร รับได้ทั้งชื่อสี (เช่น "red", "lightblue")
                                   หรือ hex (เช่น "FF0000")
        - cell_fill             : สีพื้นหลัง cell รับได้ทั้งชื่อสีหรือ hex เหมือน color
        - align_horizontal      : "left" / "center" / "right"
        - align_vertical        : "top" / "center" / "bottom"
        - number_format         : รูปแบบการแสดงผลตัวเลข เช่น "#,##0.00", "h:mm AM/PM", "dd/mm/yyyy"
        - border                : True เพื่อใส่เส้นขอบบางรอบ cell (ส่วนเสริมเพิ่มจากของเดิม)

        ตัวอย่าง:
            Set Styles    A1:D4    bold=${TRUE}    cell_fill=lightblue    align_horizontal=center    number_format=h:mm AM/PM
            Set Styles    E2    strikethrough=${TRUE}    color=FF0000
        """
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        self._check_workbook_opened()
        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet

        font_color = self._resolve_color(color)
        fill_color = self._resolve_color(cell_fill)

        font = Font(
            name=font_name,
            size=size,
            bold=bool(bold),
            italic=bool(italic),
            strike=bool(strikethrough),
            underline="single" if underline else None,
            color=font_color,
        )
        fill = (
            PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
            if fill_color
            else None
        )
        alignment = (
            Alignment(horizontal=align_horizontal, vertical=align_vertical)
            if (align_horizontal or align_vertical)
            else None
        )
        cell_border = None
        if border:
            side = Side(style="thin")
            cell_border = Border(left=side, right=side, top=side, bottom=side)

        for cell in self._iter_cells(sheet, cell_range):
            cell.font = font
            if fill:
                cell.fill = fill
            if alignment:
                cell.alignment = alignment
            if cell_border:
                cell.border = cell_border
            if number_format:
                cell.number_format = number_format

        logger.info(f"ตั้งค่า style ให้ {cell_range} แล้ว (sheet: {sheet.title})")

    # ------------------------------------------------------------------
    # รูปภาพ
    # ------------------------------------------------------------------

    @keyword("Insert Image To Worksheet")
    def insert_image_to_worksheet(self, cell, image_path, sheet_name=None, width=None, height=None):
        """แทรกรูปภาพลงใน worksheet โดยยึดมุมบนซ้ายไว้ที่ cell ที่ระบุ

        - cell       : cell ที่จะวางมุมบนซ้ายของรูป เช่น "A1"
        - image_path : path ของไฟล์รูปภาพ
        - sheet_name : sheet ที่จะแทรก (ไม่ระบุ = active worksheet)
        - width, height : (ไม่บังคับ) กำหนดขนาดรูปเป็นพิกเซล ถ้าไม่ระบุจะใช้ขนาดต้นฉบับ
        """
        from openpyxl.drawing.image import Image

        self._check_workbook_opened()
        if not os.path.isfile(image_path):
            raise FileNotFoundError(f"ไม่พบไฟล์รูปภาพ: {image_path}")

        sheet = self.workbook[sheet_name] if sheet_name else self.active_sheet
        image = Image(image_path)
        if width:
            image.width = int(width)
        if height:
            image.height = int(height)

        sheet.add_image(image, cell)
        logger.info(f"แทรกรูปภาพ '{image_path}' ที่ {cell} แล้ว (sheet: {sheet.title})")

    # ------------------------------------------------------------------
    # internal helper
    # ------------------------------------------------------------------

    def _check_workbook_opened(self):
        if self.workbook is None:
            raise RuntimeError(
                "ยังไม่ได้เปิดหรือสร้าง workbook กรุณาเรียก 'Open Workbook' หรือ 'Create Workbook' ก่อน"
            )

    @staticmethod
    def _resolve_column(column):
        """แปลงค่า column ที่รับมาให้เป็นเลข index (1-based)

        รองรับทั้ง:
        - เลข หรือ string ตัวเลข เช่น 1, "1"
        - ตัวอักษรคอลัมน์ เช่น "A", "b", "AA"
        """
        column_str = str(column).strip()
        if column_str.isdigit():
            return int(column_str)

        from openpyxl.utils import column_index_from_string

        try:
            return column_index_from_string(column_str.upper())
        except ValueError:
            raise ValueError(
                f"column ไม่ถูกต้อง: '{column}' ต้องเป็นเลข (เช่น 1, 2) หรือตัวอักษรคอลัมน์ (เช่น A, B, AA)"
            )

    @staticmethod
    def _iter_cells(sheet, cell_range):
        """คืนค่า cell object ทั้งหมดใน cell เดียวหรือช่วง cell แบบ flatten เป็น list เดียว"""
        target = sheet[cell_range]
        if isinstance(target, tuple):
            return [cell for row in target for cell in row]
        return [target]

    _COLOR_NAMES = {
        "black": "000000",
        "white": "FFFFFF",
        "red": "FF0000",
        "green": "008000",
        "blue": "0000FF",
        "yellow": "FFFF00",
        "orange": "FFA500",
        "purple": "800080",
        "pink": "FFC0CB",
        "cyan": "00FFFF",
        "magenta": "FF00FF",
        "gray": "808080",
        "grey": "808080",
        "lightblue": "ADD8E6",
        "lightgreen": "90EE90",
        "lightyellow": "FFFFE0",
        "lightgray": "D3D3D3",
        "lightgrey": "D3D3D3",
        "darkred": "8B0000",
        "darkblue": "00008B",
        "darkgreen": "006400",
        "brown": "A52A2A",
    }

    @classmethod
    def _resolve_color(cls, color):
        """แปลงค่า color ที่รับมา (ชื่อสี หรือ hex) ให้เป็น hex string ที่ openpyxl ใช้ได้

        รองรับ:
        - ชื่อสีที่รู้จัก เช่น "red", "lightblue" (ไม่สนตัวพิมพ์เล็ก/ใหญ่)
        - hex เช่น "FF0000" หรือ "#FF0000" (6 หรือ 8 หลัก)
        - None (ไม่ตั้งค่าสี)
        """
        if not color:
            return None

        color_str = str(color).strip().lstrip("#")
        lowered = color_str.lower()
        if lowered in cls._COLOR_NAMES:
            return cls._COLOR_NAMES[lowered]

        if re.fullmatch(r"[0-9A-Fa-f]{6}", color_str) or re.fullmatch(r"[0-9A-Fa-f]{8}", color_str):
            return color_str.upper()

        raise ValueError(
            f"สีไม่ถูกต้อง: '{color}' ต้องเป็นชื่อสีที่รู้จัก (เช่น red, lightblue) "
            f"หรือ hex (เช่น FF0000)"
        )
